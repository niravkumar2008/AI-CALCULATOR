"""Builds hardware/production/COST_MODEL.xlsx (formulas, blue inputs, sources beside them)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"C:\Users\r_kas\OneDrive\Documents\GitHub\AI-CALCULATOR\hardware\production\COST_MODEL.xlsx"
# LCD panel price at 100 / 1,000. Placeholder until an Alibaba quote exists (hardware/production/lcd_suppliers.json).
LCD_BULK = [2.50, 2.30]
LCD_BULK_SRC = ("hardware/production/lcd_suppliers.json (2026-10-08): Alibaba, Zhengzhou Zhongjingyuan (ZJY) ST7789V3 30-pin, $2.50 @100, $2.30 @1,000, MOQ 2 (prices from search snippets; confirm drawing + pin table and get a written quote). Sample: Shenzhen Goldenmorning T190X7-C30-01 $1.85-2.50. No listing gave a short-tail FPC price.")
LCD100_ALIBABA = 42.22  # Alibaba_Bulk_Sourcing.xlsx 'v15-LCD cost by batch' M23, 2026-10-08 (42.2155 after the 10/8-evening R23 update)

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=13, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="555555")
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
TOT_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF9E5")
thin = Side(style="thin", color="BBBBBB")
TOP = Border(top=Side(style="thin", color="000000"))
USD = '"$"#,##0.00'
USD0 = '"$"#,##0'
PCT = '0.0%'
NUM = '#,##0'
NUM2 = '#,##0.00'
WRAP = Alignment(wrap_text=True, vertical="top")

wb = openpyxl.Workbook()
names = {}  # name -> "'Sheet'!$X$n"


def ref(sheet, col, row, absolute=True):
    c = f"${col}${row}" if absolute else f"{col}{row}"
    return f"'{sheet}'!{c}"


def put(ws, cell, value, font=BLACK, fmt=None, fill=None, bold=False):
    c = ws[cell]
    c.value = value
    c.font = Font(name=F, size=10, bold=bold, color=font.color) if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = BOLD
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[L(i)].width = w


# ---------------------------------------------------------------- Inputs
wsI = wb.active
wsI.title = "Inputs"
S = "Inputs"
put(wsI, "A1", "Inputs: every blue number is an assumption you can edit. Black = formula.", TITLE)
put(wsI, "A2", "Written 2026-10-06, updated 2026-10-08 evening (5 v14 testers, v15-LCD verified ORDER with the BuyDisplay ER-TFT019-1 panel, bigger battery). Sources beside each input. (estimate) = no hard source; confirm by quote.", NOTE)
header(wsI, 4, ["Input", "Value", "Unit", "Source / why"], [46, 12, 14, 110])

inputs = [
    ("sec", "Pricing and fees"),
    ("hw_price", "Calculator price", 225, "$", USD, "Decided 2026-10-04 (LAUNCH_ROADMAP.md §0)."),
    ("sub_price", "Subscription price per month", 15, "$/month", USD, "Decided 2026-10-04 (LAUNCH_ROADMAP.md §0)."),
    ("free_months", "Free months at the start", 1, "months", NUM, "First month free = proxy TRIAL_DAYS=30 (server/proxy/config.py)."),
    ("monthly_cap", "Fair-use cap, solves per month", 500, "solves", NUM, "Proxy MONTHLY_CAP=500 (server/proxy/config.py)."),
    ("pay_pct", "Card fee, percent", 0.029, "%", PCT, "Stripe US standard online card rate 2.9 % + 30¢ (estimate; check your Stripe dashboard)."),
    ("pay_fixed", "Card fee, fixed per charge", 0.30, "$", USD, "Stripe US standard 30¢ per charge (estimate)."),
    ("ship_out", "Shipping one calculator to a customer", 6.00, "$", USD, "USPS Ground Advantage, ~8 oz padded box, zone 4-5 (estimate)."),
    ("warranty_pct", "Returns / warranty allowance (% of price)", 0.03, "%", PCT, "Roadmap §5: budget 2-5 % of units needing help or replacement (estimate)."),
    ("sec", "Subscription behaviour"),
    ("attach", "Buyers who start paying after the free month", 0.70, "%", PCT, "Assumption (estimate). Nothing measured yet: replace with beta data."),
    ("churn", "Monthly cancellation rate of paying users", 0.08, "%/month", PCT, "Assumption (estimate). Student products churn in summer; 8 %/month ≈ half gone after 8 months."),
    ("horizon", "Months counted for subscription value (LTV)", 12, "months", NUM, "One school year plus summer. Conservative; anything after month 12 is upside."),
    ("server_month", "Proxy hosting per month (server + database)", 25, "$/month", USD, "Small Render/Fly instance + disk (estimate). server/proxy/README.md: not deployed yet."),
    ("sec", "Claude API usage (per solve)"),
    ("model", "Model used for solves", "Opus 5.5", "", None, "Pick from the table below (Opus 5.5 / Sonnet 5.5 / Haiku 4.5). NOTE: the proxy's default MAIN_MODEL today is claude-sonnet-5-5 (config.py)."),
    ("solves_day", "Solves per active user per day (typical)", 5, "solves/day", NUM2, "Roadmap §3.3 'typical' = 150 solves/month ≈ 5/day (estimate)."),
    ("days_month", "Days per month", 30.4, "days", NUM2, "365 / 12."),
    ("in_tokens", "Input tokens per solve (photo + prompt)", 2500, "tokens", NUM, "Roadmap §3.3: photo ≈ (w×h)/750 ≈ 1,600 for 1280×960, plus system prompt ≈ 2,500 total (estimate)."),
    ("cached_tokens", "…of which a fixed system prompt that can be cached", 900, "tokens", NUM, "Estimate. If the fixed prompt is shorter than the model's minimum cacheable prefix (512-4,096 tokens) it will not cache: set 0."),
    ("cache_hit", "Cache hit rate on that prefix", 0.80, "%", PCT, "Estimate (cache lives ~5 min; busy hours hit, quiet hours miss)."),
    ("cache_write_mult", "Cache write price as a multiple of input price", 1.25, "×", NUM2, "Anthropic 5-minute cache write = 1.25× base input."),
    ("batch_disc", "Batch discount applied", 0.0, "%", PCT, "Batch API is 50 % off but answers take minutes to hours: unusable for live solves. Keep 0 %."),
    ("sec", "Selling legally (one-time, before public sales)"),
    ("fcc", "FCC Part 15B SDoC test (1 lab day)", 1750, "$", USD0, "Roadmap §4: $1,000-2,500 [S10]; midpoint."),
    ("lawyer", "Lawyer: trade dress + terms of service", 300, "$", USD0, "Roadmap §4: clinic $0-300; KNOCKOFF_SHELL_PLAN.md §5: private consult $200-500."),
    ("llc", "Indiana LLC", 97, "$", USD0, "Roadmap §4 [S25]."),
    ("trademark", "USPTO trademark, classes 9 + 42", 700, "$", USD0, "Roadmap §4: $350 per class [S24]."),
    ("insurance", "Product liability insurance, first year", 700, "$", USD0, "Roadmap §4: ~$400-1,000/yr (estimate)."),
    ("sec", "JLCPCB assembly fee check (already inside the PCBA price row)"),
    ("jlc_setup", "Standard PCBA setup fee per order", 25, "$", USD, "Given 2026-10-06 (JLCPCB Standard PCBA)."),
    ("ext_fee", "Extended-part loading fee per part type", 3, "$", USD, "Roadmap §3.1: $3 × 14 extended types per order (JLCPCB rule)."),
    ("ext_types", "Extended part types on the board", 14, "types", NUM, "Roadmap §3.1."),
    ("joints", "Solder joints per board", 310, "joints", NUM, "Estimate from the 77-designator BOM: ESP32 module 65, TCA8418 25, 2 FPCs 52, SOT-23s ~45, connectors ~12, ~55 passives × 2."),
    ("joint_rate", "Price per solder joint at 100 boards", 0.0013, "$", '"$"0.0000', "Given 2026-10-06 (JLCPCB)."),
    ("sec", "Your own time (shown, not added to cost at 2-100)"),
    ("own_min", "Your minutes per unit: grind + assemble + QA", 37, "min", NUM, "Grinds with jigs 9-11 min (jigs/README.md), assembly ~15 min (estimate), QA Part A 5 + Part B 6-8 min (QA_TEST_PLAN.md)."),
    ("own_rate", "Value of your hour", 15, "$/h", USD, "Assumption: a student job wage (estimate)."),
    ("sec", "v15-LCD board change (per board; hardware/stage15_lcd.md §3, verification/13 and 14)"),
    ("v15_removed", "Parts removed: 17 e-paper parts (J2, L1, Q3, R11, R12, D3-D5, C20-C30)", 0.75, "$/board", USD, "Estimate from LCSC 1-off prices: J2 FPC socket ~$0.30, L1 68 µH ~$0.12, Q3 ~$0.15, 3 × B5819W ~$0.06, C20 + 10 × 1 µF/50 V ~$0.11, 2 resistors. Stage 15: board BOM ≈ $1 cheaper."),
    ("v15_added", "Parts added: J5 C2919501 $0.36, Q4 AO3401A $0.10, Q5 AO3400A C20917 $0.09, R21/R22/C36, R23 15 Ω C22810", 0.5573, "$/board", USD, "JLCPCB prices 2026-10-08 (stage15_lcd.md §3 table): 0.36 + 0.10 + 0.09 + 3 × 0.002 + R23 C22810 0.0013 (review 14: 22 Ω → 15 Ω)."),
    ("v15_ext_delta", "Change in JLCPCB extended part types (v15 vs v14)", -1, "types", NUM, "Review 14 stock list: v14 had 14 extended types, v15 has 13. L1 (C135265), Q3 (C469327) and R12 (C23157) leave; J5 (C2919501) and R23 15 Ω (C22810, ~$3 fee) join; J1 keeps C6364666."),
    ("v15_order", "v15-LCD JLCPCB order: 5 PCBs, 5 assembled, delivered (DHL + ~35 % duty)", 270, "$", USD0, "Estimate (hardware/ORDER_WALKTHROUGH_v15_lcd.md §8): the real v14 total $275 minus one extended-part fee and the ~$0.19/board cheaper BOM, with duty; range ≈ $255-285. Replace with the real quote."),
    ("v15_reuse", "Assembled v15 boards from that order that go into the pilot", 5, "boards", NUM, "Plan: the 5 bench-tested v15 boards become pilot units 1-5, so the pilot pays for fewer new boards. Set 0 to keep them as spares."),
    ("sec", "Budget"),
    ("budget", "Cash available for supply", 2000, "$", USD0, "Business plan: $2,000 budget."),
    ("testers", "v14 e-paper testers being built", 5, "units", NUM, "Plan 2026-10-08: 5 testers in Casio shells."),
    ("pilot_units", "v15-LCD pilot size (knockoff shells)", 10, "units", NUM, "Plan: after the v15 order. Lower it (e.g. 5) if the cash budget goes negative."),
]

r = 5
for item in inputs:
    if item[0] == "sec":
        put(wsI, f"A{r}", item[1], BOLD)
        r += 1
        continue
    key, label, val, unit, fmt, src = item
    put(wsI, f"A{r}", label)
    c = put(wsI, f"B{r}", val, BLUE, fmt, INPUT_FILL)
    put(wsI, f"C{r}", unit)
    d = put(wsI, f"D{r}", src)
    d.alignment = WRAP
    names[key] = ref(S, "B", r)
    r += 1

# model price table
r += 1
put(wsI, f"A{r}", "Claude model prices (per million tokens)", BOLD)
r += 1
for i, lab in enumerate(["Model", "Input $/M", "Output $/M", "Cache read $/M", "Output tokens per solve", "Source"], start=1):
    c = wsI.cell(row=r, column=i, value=lab)
    c.font = BOLD
    c.fill = HDR_FILL
# widen E/F for table
wsI.column_dimensions["E"].width = 14
wsI.column_dimensions["F"].width = 70
mt_first = r + 1
models = [
    ("Opus 5.5", 4, 20, 0.20, 2100, "Given 2026-10-06 and Anthropic price table (cached 2026-09-25). Output incl. thinking (roadmap §3.3)."),
    ("Sonnet 5.5", 2, 10, 0.20, 2100, "Anthropic price table (cached 2026-09-25). Proxy default MAIN_MODEL."),
    ("Haiku 4.5", 1, 5, 0.10, 800, "Anthropic price table; cache read = 0.1× input (estimate). Roadmap §3.3: ~800 output tokens."),
]
for m in models:
    r += 1
    put(wsI, f"A{r}", m[0], BLUE, None, INPUT_FILL)
    put(wsI, f"B{r}", m[1], BLUE, USD, INPUT_FILL)
    put(wsI, f"C{r}", m[2], BLUE, USD, INPUT_FILL)
    put(wsI, f"D{r}", m[3], BLUE, USD, INPUT_FILL)
    put(wsI, f"E{r}", m[4], BLUE, NUM, INPUT_FILL)
    put(wsI, f"F{r}", m[5]).alignment = WRAP
mt_last = r
MT = {k: f"'{S}'!${c}${mt_first}:${c}${mt_last}" for k, c in
      [("name", "A"), ("in", "B"), ("out", "C"), ("cache", "D"), ("otok", "E")]}
dv = DataValidation(type="list", formula1=f"=$A${mt_first}:$A${mt_last}", allow_blank=False)
wsI.add_data_validation(dv)
dv.add(names["model"].split("!")[1].replace("$", ""))

r += 2
put(wsI, f"A{r}", "Derived certification total (formula)", BOLD)
put(wsI, f"B{r}", f"={names['fcc']}+{names['lawyer']}+{names['llc']}+{names['trademark']}+{names['insurance']}", BLACK, USD0)
names["cert_total"] = ref(S, "B", r)
put(wsI, f"D{r}", "FCC + lawyer + LLC + trademark + insurance. Roadmap §3 calls this 'about $3.5k one-time'.")
wsI.freeze_panes = "A5"

# ---------------------------------------------------------------- Unit cost
wsU = wb.create_sheet("Unit cost")
U = "Unit cost"
put(wsU, "A1", "Cost per finished calculator at 5 / 10 / 100 / 1,000 units (e-paper version; LCD and big battery on 'Versions')", TITLE)
put(wsU, "A2", "5 = the v14 e-paper TESTERS in genuine Casio fx-115ES shells (option C of Claude outputs/AI_Calculator_Shopping_List.xlsx). 10 / 100 / 1,000 = knockoff 991ES-style shell (v15 board). "
               "Per-unit prices from Claude outputs/Alibaba_Bulk_Sourcing.xlsx 'Cost by batch' and LAUNCH_ROADMAP.md §3 unless noted.", NOTE)
header(wsU, 4, ["#", "Item", "5 testers", "10 units", "100 units", "1,000 units", "Source / notes"],
       [5, 52, 13, 13, 13, 13, 95])
VOLS = ["C", "D", "E", "F"]
put(wsU, "B5", "Batch size (units)", BOLD)
for col, v in zip(VOLS, [5, 10, 100, 1000]):
    put(wsU, f"{col}5", v, BLUE, NUM, INPUT_FILL)
put(wsU, "G5", "Edit to test other batch sizes; the per-unit prices below don't re-scale by themselves.")
put(wsU, "B6", "Shell used")
for col, v in zip(VOLS, ["Casio donor", "Clone (retail)", "Clone (Alibaba)", "Clone (factory)"]):
    put(wsU, f"{col}6", v, BLUE, None, INPUT_FILL)

put(wsU, "A7", "Per-unit parts (price per unit at that volume)", BOLD)
rows_var = [
    ("PCB + PCBA (board, all SMD parts, JLCPCB assembly)", [55.00, 18, 9, 7.30],
     "5: JLCPCB v14 DELIVERED (incl. DHL + ~35 % duty) ≈ $220 for 2 assembled + $55 for 5 assembled = $275 ÷ 5 (Shopping List rows 1-2). 10/100/1,000: roadmap §3 table [S8]."),
    ("Shell + key mat + window + screws", [20, 6, 3.99, 3.25],
     "5: Casio fx-115ES $20 each (3 bought, 2 already owned; value shown for all 5). 10: retail clone incl. shipping. 100: Alibaba $3.25-3.99 MOQ 20 [S1]. 1,000: $3.25 (factories quote $1.50-3.50 at 3,000 [S2])."),
    ("Camera OV5640 AF", [10.594, 12, 6, 4.50], "5: option C = (3 × Seeed 114993115 $12.99 + 2 × generic $7) ÷ 5. Then roadmap / Alibaba OEM [S3]."),
    ("E-paper 2.13\" raw panel", [5.994, 6.60, 5.50, 4.00], "5: option C = (3 × Waveshare $6.99 + 2 × GDEY0213B74 $4.50) ÷ 5. Then Good Display direct by RFQ (estimate)."),
    ("LiPo 150 mAh + JST-PH", [4.77, 5.95, 1.80, 1.30], "5: option C = (3 × Adafruit #1317 $5.95 + 2 × generic 302030 $3) ÷ 5. 10: Adafruit #1317. 100/1,000: Alibaba cell with PH lead + PCM (estimate)."),
    ("Battery dangerous-goods freight surcharge", [0, 0, 0.80, 0.30], "UN3481 air surcharge (estimate)."),
    ("Magnetic connector pair", [6.50, 5.85, 1.80, 1.20], "Adafruit #5358 $6.50 / $5.85 (10+); Alibaba MG04-254-RA $1.80 [S7]."),
    ("Magnetic USB cable (in the box)", [4.95, 4.95, 1.80, 1.40], "Adafruit #5412 $4.95; OEM $1.40-1.80 (estimate)."),
    ("Retail box + insert", [0, 1.50, 1.00, 0.60], "Prototypes unboxed. Alibaba rigid box $0.20-1.02 @100 [S9]."),
    ("Quick-start card (with 'not affiliated with Casio')", [0, 0.10, 0.25, 0.05], "[S10]; small runs estimate."),
    ("Serial / QC label + polybag", [0, 0.15, 0.10, 0.05], "Estimate."),
    ("Tape, Kapton, glue (per-unit share)", [0, 0.50, 0.30, 0.20], "2: bought as bench tools (one-time row below). Estimate."),
    ("Inbound freight + US import duty (per unit)", [0, 8, 8, 5], "5: JLCPCB freight/duty is inside row 1; other shipping is a lump (one-time row below). Roadmap: ~35 % duty, low end of 35-92.5 % [S8]. Biggest swing factor."),
]
r = 8
var_first = r
for i, (lab, vals, src) in enumerate(rows_var, start=1):
    put(wsU, f"A{r}", i)
    put(wsU, f"B{r}", lab)
    for col, v in zip(VOLS, vals):
        put(wsU, f"{col}{r}", v, BLUE, USD, INPUT_FILL)
    put(wsU, f"G{r}", src).alignment = WRAP
    r += 1
var_last = r - 1
put(wsU, f"B{r}", "Parts subtotal", BOLD)
for col in VOLS:
    put(wsU, f"{col}{r}", f"=SUM({col}{var_first}:{col}{var_last})", BLACK, USD).fill = TOT_FILL
sub_row = r
r += 1
put(wsU, f"B{r}", "Scrap + spares allowance, %")
for col, v in zip(VOLS, [0, 0.10, 0.07, 0.05]):
    put(wsU, f"{col}{r}", v, BLUE, PCT, INPUT_FILL)
put(wsU, f"G{r}", "Alibaba sheet row 15 (estimate). At 5 the spares are a one-time row instead.")
scrap_pct = r
r += 1
put(wsU, f"B{r}", "Scrap + spares allowance, $")
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{sub_row}*{col}{scrap_pct}", BLACK, USD)
scrap_row = r
r += 1
put(wsU, f"B{r}", "Paid assembly labour (box build)")
for col, v in zip(VOLS, [0, 0, 0, 2.50]):
    put(wsU, f"{col}{r}", v, BLUE, USD, INPUT_FILL)
put(wsU, f"G{r}", "You build 5-100 yourself (time not costed, see row below). 1,000: EMS box build (roadmap §3, estimate).")
lab_row = r
r += 1
put(wsU, f"B{r}", "Variable cost per unit", BOLD)
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{sub_row}+{col}{scrap_row}+{col}{lab_row}", BLACK, USD, TOT_FILL, bold=True)
varc_row = r
r += 2

put(wsU, f"A{r}", "One-time costs for the batch (whole $ amount)", BOLD)
r += 1
rows_once = [
    ("Other shipping lump (Adafruit $15 + Seeed $12 + Waveshare $20 + AliExpress $10; JLCPCB's is in row 1)", [57, 0, 0, 0],
     "Shopping List (option C) shipping rows. At 10+ freight is per unit (row 13)."),
    ("Bench tools and consumables (tape, glue, drill bit, vinyl, IPA, glasses)", [61, 0, 0, 0], "Shopping List supplies section (charger qty 0 = you own one)."),
    ("Tester spares (option C: 1 spare of each current + generic camera, screen, battery; 1 magnet piece)", [46.93, 0, 0, 0], "12.99 + 7 + 6.99 + 4.50 + 5.95 + 3 + 6.50 (Shopping List qty 6 of each kind)."),
    ("Supplier samples (2 each from 2 suppliers, incl. 4 knockoff shells) + express", [0, 150, 150, 120], "Alibaba sheet row 16 (estimate); KNOCKOFF_SHELL_PLAN.md §1."),
    ("Pogo test jig + fixtures", [0, 0, 100, 150], "Roadmap §5: $50-150 (estimate)."),
    ("Third-party pre-shipment inspection", [0, 0, 0, 300], "~1 man-day (estimate). Only worth it at 1,000."),
]
once_first = r
for lab, vals, src in rows_once:
    put(wsU, f"B{r}", lab).alignment = WRAP
    for col, v in zip(VOLS, vals):
        put(wsU, f"{col}{r}", v, BLUE, USD, INPUT_FILL)
    put(wsU, f"G{r}", src).alignment = WRAP
    r += 1
once_last = r - 1
put(wsU, f"B{r}", "One-time total for the batch", BOLD)
for col in VOLS:
    put(wsU, f"{col}{r}", f"=SUM({col}{once_first}:{col}{once_last})", BLACK, USD, TOT_FILL)
once_tot = r
r += 1
put(wsU, f"B{r}", "One-time per unit")
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{once_tot}/{col}$5", BLACK, USD)
once_pu = r
r += 2

put(wsU, f"B{r}", "COST PER UNIT (before selling-legally costs)", BOLD)
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{varc_row}+{col}{once_pu}", BLACK, USD, TOT_FILL, bold=True)
put(wsU, f"G{r}", "Compare: Alibaba_Bulk_Sourcing.xlsx ≈ $92 @10, $46 @100, $34 @1,000 (see 'Checks').")
cpu_row = r
names["cpu"] = {col: ref(U, col, r) for col in VOLS}
r += 1
put(wsU, f"B{r}", "Selling legally? (1 = yes, spread the one-time certification over this batch)")
for col, v in zip(VOLS, [0, 0, 1, 1]):
    put(wsU, f"{col}{r}", v, BLUE, NUM, INPUT_FILL)
put(wsU, f"G{r}", "5 = testers (not sold); 10 = free / at-cost beta (KNOCKOFF_SHELL_PLAN.md §5). Public sales start at 100.")
sell_row = r
r += 1
put(wsU, f"B{r}", "Certification, lawyer, LLC, trademark, insurance per unit")
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{sell_row}*{names['cert_total']}/{col}$5", BLACK, USD)
cert_pu = r
r += 1
put(wsU, f"B{r}", "COST PER UNIT incl. selling-legally costs", BOLD)
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{cpu_row}+{col}{cert_pu}", BLACK, USD, TOT_FILL, bold=True)
cpu_all = r
names["cpu_all"] = {col: ref(U, col, r) for col in VOLS}
r += 1
put(wsU, f"B{r}", "Cash for the whole batch (incl. one-time)", BOLD)
for col in VOLS:
    put(wsU, f"{col}{r}", f"={col}{cpu_all}*{col}$5", BLACK, USD0)
batch_cash = r
names["batch_cash"] = {col: ref(U, col, r) for col in VOLS}
r += 2

put(wsU, f"A{r}", "Information only (not added above)", BOLD)
r += 1
put(wsU, f"B{r}", "JLCPCB assembly fee inside the PCBA price: (setup + extended fees) ÷ batch + joints × rate")
for col in VOLS:
    put(wsU, f"{col}{r}",
        f"=({names['jlc_setup']}+{names['ext_fee']}*{names['ext_types']})/{col}$5+{names['joints']}*{names['joint_rate']}",
        BLACK, USD)
put(wsU, f"G{r}", "Why the board price drops from 10 to 100: the fixed $67 per order is shared. The joint fee itself is only ~$0.40 a board; the rest of row 1 is parts.")
r += 1
put(wsU, f"B{r}", "Your time per unit, valued (not cash)")
for col in VOLS:
    put(wsU, f"{col}{r}", f"=IF({col}{lab_row}>0,0,{names['own_min']}/60*{names['own_rate']})", BLACK, USD)
put(wsU, f"G{r}", "37 min × $15/h. Zero where paid labour takes over.")
names["once_tot_row"] = once_tot
names["scrap_pct_row"] = scrap_pct
names["varc_row"] = varc_row
names["sell_row"] = sell_row
names["var_first"] = var_first
wsU.freeze_panes = "C5"

# ---------------------------------------------------------------- API cost
wsA = wb.create_sheet("API cost")
A = "API cost"
put(wsA, "A1", "Claude API cost per solve and per subscriber per month", TITLE)
put(wsA, "A2", "Chosen model and usage come from 'Inputs'. Batch pricing does not apply to live solves.", NOTE)
header(wsA, 4, ["Item", "Value", "How it's worked out"], [52, 14, 90])
a_rows = [
    ("Model", f"={names['model']}", None, "Inputs"),
    ("Input price $/M tokens", f"=INDEX({MT['in']},MATCH(B5,{MT['name']},0))", USD, "Model table on Inputs"),
    ("Output price $/M tokens", f"=INDEX({MT['out']},MATCH(B5,{MT['name']},0))", USD, "Model table"),
    ("Cache read price $/M tokens", f"=INDEX({MT['cache']},MATCH(B5,{MT['name']},0))", USD, "Model table"),
    ("Output tokens per solve (incl. thinking)", f"=INDEX({MT['otok']},MATCH(B5,{MT['name']},0))", NUM, "Model table"),
    ("Input tokens read from cache", f"={names['cached_tokens']}*{names['cache_hit']}", NUM, "cached prefix × hit rate"),
    ("Input tokens written to cache (misses)", f"={names['cached_tokens']}*(1-{names['cache_hit']})", NUM, "cached prefix × miss rate"),
    ("Input tokens at full price", f"={names['in_tokens']}-B10", NUM, "all input − cache reads (writes are billed at full price plus the write premium)"),
    ("Input cost per solve", f"=(B12*B6+B10*B8+B11*B6*({names['cache_write_mult']}-1))/1000000", '"$"0.0000', "full-price input + cache reads + write premium"),
    ("Output cost per solve", f"=B9*B7/1000000", '"$"0.0000', "output tokens × output price"),
    ("COST PER SOLVE", f"=(B13+B14)*(1-{names['batch_disc']})", '"$"0.0000', "after any batch discount"),
    ("Solves per user per month (typical)", f"={names['solves_day']}*{names['days_month']}", NUM, "solves/day × days"),
    ("API COST PER SUBSCRIBER PER MONTH (typical)", f"=B15*B16", USD, ""),
    ("API cost per month for a user at the fair-use cap", f"=B15*{names['monthly_cap']}", USD, "cap = 500 solves (proxy)"),
    ("Solves per month that $15 can pay for", f"=({names['sub_price']}-({names['sub_price']}*{names['pay_pct']}+{names['pay_fixed']}))/B15", NUM, "($15 − card fee) ÷ cost per solve. Hosting excluded."),
    ("…as solves per day", f"=B19/{names['days_month']}", NUM2, ""),
    ("Cap that keeps a heavy user at break-even (suggested MONTHLY_CAP)", f"=ROUNDDOWN(B19,-1)", NUM, "rounded down to 10"),
]
for i, (lab, f_, fmt, how) in enumerate(a_rows, start=5):
    put(wsA, f"A{i}", lab, BOLD if lab.isupper() or lab.startswith("API COST") or lab.startswith("COST") else BLACK)
    c = put(wsA, f"B{i}", f_, BLACK, fmt)
    if lab.startswith("COST") or lab.startswith("API COST"):
        c.fill = TOT_FILL
        c.font = BOLD
    put(wsA, f"C{i}", how)
names["api_month"] = ref(A, "B", 17)
names["api_solve"] = ref(A, "B", 15)
names["api_cap"] = ref(A, "B", 18)
names["api_be_day"] = ref(A, "B", 20)
names["api_be_cap"] = ref(A, "B", 21)

# scenario table
r = 28
put(wsA, f"A{r}", "Scenarios: API cost per user per month (same cache assumptions)", BOLD)
r += 1
for i, lab in enumerate(["Model", "Cost per solve", "Light 2/day", "Typical (Inputs)", "Heavy 10/day", "At the cap"], start=1):
    c = wsA.cell(row=r, column=i, value=lab)
    c.font = BOLD
    c.fill = HDR_FILL
for col in "DEF":
    wsA.column_dimensions[col].width = 15
wsA.column_dimensions["C"].width = 90
put(wsA, f"G{r}", "Light / heavy solves per day:", BOLD)
put(wsA, f"H{r}", 2, BLUE, NUM, INPUT_FILL)
put(wsA, f"I{r}", 10, BLUE, NUM, INPUT_FILL)
light, heavy = f"$H${r}", f"$I${r}"
for k, m in enumerate(models):
    rr = r + 1 + k
    mrow = mt_first + k
    pin, pout, pc, ot = (f"'{S}'!$B${mrow}", f"'{S}'!$C${mrow}", f"'{S}'!$D${mrow}", f"'{S}'!$E${mrow}")
    cr = f"{names['cached_tokens']}*{names['cache_hit']}"
    cw = f"{names['cached_tokens']}*(1-{names['cache_hit']})"
    put(wsA, f"A{rr}", f"='{S}'!$A${mrow}")
    put(wsA, f"B{rr}", f"=(({names['in_tokens']}-{cr})*{pin}+{cr}*{pc}+{cw}*{pin}*({names['cache_write_mult']}-1)+{ot}*{pout})/1000000*(1-{names['batch_disc']})", BLACK, '"$"0.0000')
    put(wsA, f"C{rr}", f"=B{rr}*{light}*{names['days_month']}", BLACK, USD)
    put(wsA, f"D{rr}", f"=B{rr}*{names['solves_day']}*{names['days_month']}", BLACK, USD)
    put(wsA, f"E{rr}", f"=B{rr}*{heavy}*{names['days_month']}", BLACK, USD)
    put(wsA, f"F{rr}", f"=B{rr}*{names['monthly_cap']}", BLACK, USD)
names["scen_first"] = r + 1
put(wsA, f"A{r+5}", "Column C is narrow-looking because it also holds the notes above; numbers are $ per user per month.", NOTE)

# ---------------------------------------------------------------- Margins & payback
wsM = wb.create_sheet("Margins & payback")
M = "Margins & payback"
put(wsM, "A1", "Margin on the $225 calculator, margin on the $15 subscription, break-even and payback", TITLE)
put(wsM, "A2", "Columns = batch size. Unit cost includes the selling-legally costs where 'Selling legally' = 1 on 'Unit cost'.", NOTE)
header(wsM, 4, ["Item", "5 testers", "10 units", "100 units", "1,000 units", "How it's worked out"], [58, 13, 13, 13, 13, 80])
MC = ["B", "C", "D", "E"]
rows_m = []


def mrow(label, fn, fmt=USD, how="", bold=False):
    rows_m.append((label, fn, fmt, how, bold))


mrow("Batch size", lambda uc, i: f"='{U}'!{uc}$5", NUM, "")
mrow("HARDWARE", None)
mrow("Price", lambda uc, i: f"={names['hw_price']}", USD, "Inputs")
mrow("Card fee", lambda uc, i: f"={names['hw_price']}*{names['pay_pct']}+{names['pay_fixed']}", USD, "2.9 % + 30¢")
mrow("Shipping to the customer", lambda uc, i: f"={names['ship_out']}", USD, "Inputs")
mrow("Returns / warranty allowance", lambda uc, i: f"={names['hw_price']}*{names['warranty_pct']}", USD, "% of price")
mrow("API cost of the free first month", lambda uc, i: f"={names['api_month']}*{names['free_months']}", USD, "Every buyer gets it, paying or not")
mrow("Unit cost (incl. selling-legally costs)", lambda uc, i: f"={names['cpu_all'][uc]}", USD, "'Unit cost' sheet")
mrow("HARDWARE PROFIT PER UNIT", lambda uc, i: f"={MC[i]}7-{MC[i]}8-{MC[i]}9-{MC[i]}10-{MC[i]}11-{MC[i]}12", USD, "price − everything above", True)
mrow("Hardware margin %", lambda uc, i: f"={MC[i]}13/{MC[i]}7", PCT, "profit ÷ price")
mrow("SUBSCRIPTION (per paying user per month)", None)
mrow("Price", lambda uc, i: f"={names['sub_price']}", USD, "Inputs")
mrow("Card fee", lambda uc, i: f"={names['sub_price']}*{names['pay_pct']}+{names['pay_fixed']}", USD, "the 30¢ hurts most on small charges")
mrow("Claude API (typical user)", lambda uc, i: f"={names['api_month']}", USD, "'API cost' sheet")
mrow("Hosting share", lambda uc, i: f"={names['server_month']}/MAX(1,{MC[i]}5*{names['attach']})", USD, "server $/month ÷ paying users (batch × attach rate)")
mrow("SUBSCRIPTION PROFIT PER USER PER MONTH", lambda uc, i: f"={MC[i]}16-{MC[i]}17-{MC[i]}18-{MC[i]}19", USD, "", True)
mrow("Subscription margin %", lambda uc, i: f"={MC[i]}20/{MC[i]}16", PCT, "")
mrow("Same, for a user at the 500-solve cap", lambda uc, i: f"={MC[i]}16-{MC[i]}17-{names['api_cap']}-{MC[i]}19", USD, "negative = that user costs you money")
mrow("Does $15 cover the API? (typical user)", lambda uc, i: f'=IF({MC[i]}18<{MC[i]}16-{MC[i]}17,"YES","NO")', None, "")
mrow("LIFETIME (per calculator sold)", None)
mrow("Expected paid months in the horizon", lambda uc, i: f"={names['attach']}*(1-(1-{names['churn']})^({names['horizon']}-{names['free_months']}))/{names['churn']}", NUM2, "attach × (1 − (1−churn)^(months−free)) ÷ churn")
mrow("Subscription profit over the horizon", lambda uc, i: f"={MC[i]}25*{MC[i]}20", USD, "")
mrow("TOTAL PROFIT PER CALCULATOR (hardware + subscription)", lambda uc, i: f"={MC[i]}13+{MC[i]}26", USD, "", True)
mrow("BREAK-EVEN AND PAYBACK", None)
mrow("One-time costs this batch carries (samples, jig, QC, certification)", lambda uc, i: f"='{U}'!{uc}{once_tot}+'{U}'!{uc}{sell_row}*{names['cert_total']}", USD0, "'Unit cost' one-time + certification if selling")
mrow("Variable cost per unit", lambda uc, i: f"='{U}'!{uc}{varc_row}", USD, "")
mrow("Profit per unit before one-time costs", lambda uc, i: f"={MC[i]}7-{MC[i]}8-{MC[i]}9-{MC[i]}10-{MC[i]}11-{MC[i]}30", USD, "")
mrow("Units to sell to cover the one-time costs (hardware only)", lambda uc, i: f"=IF({MC[i]}31>0,ROUNDUP({MC[i]}29/{MC[i]}31,0),\"never\")", NUM, "one-time ÷ profit per unit")
mrow("Months for one paying user to repay the free month's API", lambda uc, i: f"=IF({MC[i]}20>0,{MC[i]}11/{MC[i]}20,\"never\")", NUM2, "free-month API ÷ subscription profit")
mrow("Months of subscription that would repay the unit cost alone", lambda uc, i: f"=IF({MC[i]}20>0,{MC[i]}12/{MC[i]}20,\"never\")", NUM2, "if the calculator were sold at cost")
mrow("Calculators to sell to get the $2,000 budget back", lambda uc, i: f"=IF({MC[i]}13>0,ROUNDUP({names['budget']}/{MC[i]}13,0),\"never\")", NUM, "$2,000 ÷ hardware profit per unit")

for i, (lab, fn, fmt, how, bold) in enumerate(rows_m, start=5):
    if fn is None:
        put(wsM, f"A{i}", lab, BOLD)
        for col in MC:
            wsM[f"{col}{i}"].fill = HDR_FILL
        wsM[f"A{i}"].fill = HDR_FILL
        continue
    put(wsM, f"A{i}", lab, BOLD if bold else BLACK)
    for k, (col, uc) in enumerate(zip(MC, VOLS)):
        c = put(wsM, f"{col}{i}", fn(uc, k), BLACK, fmt)
        if bold:
            c.font = BOLD
            c.fill = TOT_FILL
    put(wsM, f"F{i}", how)
# sanity: verify row labels align with the fixed row numbers used in formulas
expect = {7: "Price", 13: "HARDWARE PROFIT", 16: "Price", 20: "SUBSCRIPTION PROFIT", 25: "Expected paid", 26: "Subscription profit", 29: "One-time", 31: "Profit per unit"}
for rr, txt in expect.items():
    assert str(wsM[f"A{rr}"].value).startswith(txt), (rr, wsM[f"A{rr}"].value)
wsM.freeze_panes = "B5"

# ---------------------------------------------------------------- Versions (LCD, bigger battery)
wsV = wb.create_sheet("Versions")
V = "Versions"
put(wsV, "A1", "Four versions at 10 / 100 / 1,000: e-paper or LCD, 150 mAh or 1,200-1,500 mAh battery", TITLE)
put(wsV, "A2", "Starts from the e-paper cost on 'Unit cost' and adds only the differences (each difference also carries that column's scrap %). "
               "LCD = v15-LCD board (hardware/stage15_lcd.md): 1.9\" 170×320 IPS ST7789 30-pin panel instead of the 2.13\" e-paper. "
               "Big battery = Adafruit #258 1,200 mAh at 10, Alibaba 504060 ~1,500 mAh at 100+ (enclosure/final_assembly/battery_upgrade.md): needs the custom shell pocket (70 × 41 × 5.4 mm) or a knockoff with room (KNOCKOFF_SHELL_PLAN.md K-D7/K-S5).", NOTE)
header(wsV, 4, ["#", "Item", "10 units", "100 units", "1,000 units", "Source / notes"], [5, 60, 13, 13, 13, 100])
VC = ["C", "D", "E"]          # columns here
UCOL = ["D", "E", "F"]        # matching 'Unit cost' columns
MCOL = ["C", "D", "E"]        # matching 'Margins & payback' columns
vf = names["var_first"]
EP_ROW, BAT_ROW, DG_ROW = vf + 3, vf + 4, vf + 5
assert str(wsU[f"B{EP_ROW}"].value).startswith("E-paper") and str(wsU[f"B{BAT_ROW}"].value).startswith("LiPo") and "dangerous" in str(wsU[f"B{DG_ROW}"].value)
r = 5
put(wsV, f"A{r}", "Inputs that differ by version", BOLD)
r += 1
vrows = {}


def vline(key, label, vals, fmt=USD, src="", formula=False):
    global r
    put(wsV, f"B{r}", label)
    for col, v in zip(VC, vals):
        if v is None:
            continue
        if formula:
            put(wsV, f"{col}{r}", v, BLACK, fmt)
        else:
            put(wsV, f"{col}{r}", v, BLUE, fmt, INPUT_FILL)
    put(wsV, f"F{r}", src).alignment = WRAP
    vrows[key] = r
    r += 1


vline("ep", "E-paper panel (from 'Unit cost')", [f"='{U}'!{u}{EP_ROW}" for u in UCOL], src="Unit cost row 4.", formula=True)
vline("ali_price", "LCD panel price at 10: BuyDisplay ER-TFT019-1 (no touch, datasheet-verified)", [6.22, None, None],
      src="Given 2026-10-08 (verification/14, LCD_PANEL_OPTIONS.md): ~$6-7 each, $6.22 @10, $5.71 @100.")
vline("ali_ship", "BuyDisplay shipping per order", [12, None, None], src="Estimate, not quoted: check at checkout and replace.")
vline("lcd_bulk", "LCD panel price at 100 / 1,000 (Alibaba / panel maker)", [None, LCD_BULK[0], LCD_BULK[1]], src=LCD_BULK_SRC)
vline("lcd", "LCD panel per unit", [f"=C{vrows['ali_price']}+C{vrows['ali_ship']}/'{U}'!D$5", f"=D{vrows['lcd_bulk']}", f"=E{vrows['lcd_bulk']}"],
      src="10: BuyDisplay price + one shipping charge shared by the batch. 100/1,000: the line above (Alibaba after a sample check).", formula=True)
vline("pcb_parts", "v15 PCBA parts change per board (added − removed)", [f"={names['v15_added']}-{names['v15_removed']}"] * 3,
      src="Inputs: v15-LCD board change.", formula=True)
vline("pcb_ext", "v15 extended-part fee change per board", [f"={names['v15_ext_delta']}*{names['ext_fee']}/'{U}'!{u}$5" for u in UCOL],
      src="types change × $3 per order ÷ batch.", formula=True)
vline("bat", "150 mAh battery + its DG surcharge (from 'Unit cost')", [f"='{U}'!{u}{BAT_ROW}+'{U}'!{u}{DG_ROW}" for u in UCOL], src="Unit cost rows 5-6.", formula=True)
vline("big", "1,200-1,500 mAh battery", [9.95, 4.00, 3.00],
      src="10: Adafruit #258 1,200 mAh $9.95 (battery_upgrade.md). 100/1,000: Alibaba 504060 ~1,500 mAh ~$3-5 at MOQ (estimate). Must come with UN38.3 in the seller's name.")
vline("big_dg", "Its dangerous-goods freight surcharge", [0, 1.20, 0.50],
      src="Estimate: ~5.5 Wh cell, heavier than the 150 mAh one (UN3481). At 10 it ships ground from Adafruit (in the price).")
vline("big_lead", "Longer JST-PH lead / extension (the pocket is under the keypad, ~150 mm lead)", [0.50, 0.25, 0.15],
      src="Estimate. Also R2 20 k → 4.7 k on the board for a ~3 h charge (one resistor, no cost change).")
r += 1
put(wsV, f"A{r}", "Differences per unit (incl. scrap %)", BOLD)
r += 1


def dline(key, label, expr, src=""):
    global r
    put(wsV, f"B{r}", label)
    for col, u in zip(VC, UCOL):
        put(wsV, f"{col}{r}", expr(col, u), BLACK, USD)
    put(wsV, f"F{r}", src).alignment = WRAP
    vrows[key] = r
    r += 1


dline("d_lcd", "LCD instead of e-paper",
      lambda c, u: f"=({c}{vrows['lcd']}-{c}{vrows['ep']}+{c}{vrows['pcb_parts']}+{c}{vrows['pcb_ext']})*(1+'{U}'!{u}{names['scrap_pct_row']})",
      "(panel − e-paper + board change) × (1 + scrap %).")
dline("d_big", "Big battery instead of 150 mAh",
      lambda c, u: f"=({c}{vrows['big']}+{c}{vrows['big_dg']}+{c}{vrows['big_lead']}-{c}{vrows['bat']})*(1+'{U}'!{u}{names['scrap_pct_row']})",
      "(big cell + its DG + lead − small cell − its DG) × (1 + scrap %).")
r += 1
put(wsV, f"A{r}", "Results by version", BOLD)
r += 1
variants = [("v_ep", "E-paper, 150 mAh (= 'Unit cost')", ""),
            ("v_lcd", "LCD, 150 mAh (v15-LCD)", "d_lcd"),
            ("v_ep_big", "E-paper, 1,200-1,500 mAh", "d_big"),
            ("v_lcd_big", "LCD, 1,200-1,500 mAh (the long-term target)", "d_lcd+d_big")]
for i, lab in enumerate(["", "Version", "10 units", "100 units", "1,000 units"], start=1):
    c = wsV.cell(row=r, column=i, value=lab or None)
    c.font = BOLD
    c.fill = HDR_FILL
r += 1


def delta(c, spec):
    if not spec:
        return ""
    return "".join(f"+{c}{vrows[k]}" for k in spec.split("+"))


put(wsV, f"B{r}", "Cost per unit before selling-legally costs", BOLD)
r += 1
for key, lab, spec in variants:
    put(wsV, f"B{r}", lab)
    for col, u in zip(VC, UCOL):
        put(wsV, f"{col}{r}", f"={names['cpu'][u]}{delta(col, spec)}", BLACK, USD, TOT_FILL if key == "v_lcd" else None)
    vrows[key] = r
    r += 1
put(wsV, f"B{r}", "Cost per unit incl. selling-legally costs (FCC, lawyer, LLC, trademark, insurance)", BOLD)
r += 1
for key, lab, spec in variants:
    put(wsV, f"B{r}", lab)
    for col, u in zip(VC, UCOL):
        put(wsV, f"{col}{r}", f"={names['cpu_all'][u]}{delta(col, spec)}", BLACK, USD)
    vrows[key + "_all"] = r
    r += 1
put(wsV, f"B{r}", "Hardware profit per $225 unit (same fees, shipping, warranty, free-month API as 'Margins & payback')", BOLD)
r += 1
for key, lab, spec in variants:
    put(wsV, f"B{r}", lab)
    for col, m in zip(VC, MCOL):
        put(wsV, f"{col}{r}", f"='{M}'!{m}7-'{M}'!{m}8-'{M}'!{m}9-'{M}'!{m}10-'{M}'!{m}11-{col}{vrows[key + '_all']}", BLACK, USD)
    vrows[key + "_profit"] = r
    r += 1
put(wsV, f"B{r}", "Variable cost per unit (no one-time costs), for the cash budget", BOLD)
r += 1
for key, lab, spec in variants[:2]:
    put(wsV, f"B{r}", lab)
    for col, u in zip(VC, UCOL):
        put(wsV, f"{col}{r}", f"='{U}'!{u}{names['varc_row']}{delta(col, spec)}", BLACK, USD)
    vrows[key + "_var"] = r
    r += 1
r += 1
put(wsV, f"A{r}", "Battery life (stage15_lcd.md §3): LCD screen-on ≈ 80-90 mA → ~1.5 h on 150 mAh (charge every 1-2 days), 13-17 h on 1,200-1,500 mAh. "
                  "The e-paper draws almost nothing between refreshes (~3 h screen-on, weekly charging). That is why the big cell matters for the LCD version.", NOTE)
for key in ["v_ep", "v_lcd", "v_ep_big", "v_lcd_big"]:
    for suf in ["", "_all", "_profit"]:
        names["V_" + key[2:] + suf] = {u: f"'{V}'!${c}${vrows[key + suf]}" for c, u in zip(VC, UCOL)}
names["V_lcd_var10"] = f"'{V}'!$C${vrows['v_lcd_var']}"
names["V_lcd_unit"] = {u: f"'{V}'!${c}${vrows['lcd']}" for c, u in zip(VC, UCOL)}
wsV.freeze_panes = "C5"

# ---------------------------------------------------------------- Cash
wsC = wb.create_sheet("Cash budget")
C = "Cash budget"
put(wsC, "A1", "Where the $2,000 goes (supply budget), in order", TITLE)
put(wsC, "A2", "Cash actually leaving your account, in the order of the 2026-10-08 plan. The 2 Casio shells you own are not cash. v14 JLCPCB was ordered 10/8; v15-LCD is verified ORDER (review 14) and is next; nothing else is bought yet.", NOTE)
header(wsC, 4, ["#", "What", "When", "Amount", "Source / notes"], [5, 62, 18, 13, 95])
cash = [
    ("v14 JLCPCB order: 5 PCBs, 5 assembled, delivered (DHL + ~35 % duty)", "Paid 10/8", 275, True,
     "Shopping List rows 1-2: ≈ $220 for 2 assembled + ≈ $55 for 5 assembled (ranges $200-240 and $50-60). Replace with the real total."),
    ("5 testers: 3 Casio shells, cameras, screens, batteries, magnet pieces, cables, shipping, supplies (option C)", "Now, before boards land", 388.97, True,
     "AI_Calculator_Shopping_List.xlsx option C total $663.97 minus the JLCPCB order. Option B would be ≈ $92 less, option A ≈ $2 less."),
    ("v15-LCD prototype panels: 3 × BuyDisplay ER-TFT019-1 (no touch) + shipping", "Now, with the v15 order", 32.25, True,
     "Shopping List sheet 'v15-LCD prototype' rows 1-2: 3 × ~$6.75 + ~$12 shipping (estimate). Bench checks at arrival: diode test, tail length, backlight current, 3.3 V."),
    ("v15-LCD JLCPCB order (verified ORDER, review 14): 5 PCBs, 5 assembled, delivered", "Now", f"={names['v15_order']}", False,
     "Inputs: v15 order estimate. Files hardware/fab_v15_lcd/, steps hardware/ORDER_WALKTHROUGH_v15_lcd.md."),
    ("Knockoff shell samples: 4 clones, 2 sellers (KNOCKOFF_SHELL_PLAN.md §1)", "During tester feedback", 60, True,
     "Alibaba pair $40-60 incl. DHL + AliExpress/Amazon pair $10-16 (estimate)."),
    ("Other supplier samples for bulk (camera, bare LCD panels, 150 mAh + 1,200-1,500 mAh cells, magnet pair + cable)", "After shells pass", 90, True,
     "Alibaba sheet row 16 total $150 minus the shells (estimate)."),
    ("v15-LCD pilot in knockoff shells: the remaining boards + all parts, freight, scrap", "After the v15 boards pass the bench checks", f"={names['pilot_units']}*{names['V_lcd_var10']}-MIN({names['pilot_units']},{names['v15_reuse']})*'{U}'!$D${var_first}", False,
     "Pilot size (Inputs) × LCD variable cost at 10 ('Versions'), minus the boards re-used from the v15 order × the $18 board cost at 10 ('Unit cost' row 1). Includes $8/unit freight and duty."),
    ("Proxy hosting, first 3 months", "At beta", f"=3*{names['server_month']}", False, "Inputs: server $/month × 3."),
    ("Claude API for the testers' and pilot users' free month", "At beta", f"=({names['testers']}+{names['pilot_units']})*{names['api_month']}", False,
     "(testers + pilot) × 1 month at the typical rate of the model chosen on Inputs."),
    ("Lawyer / IP clinic: trade-dress check before any unit leaves the house", "Before testers go out", 150, True,
     "KNOCKOFF_SHELL_PLAN.md §5. University clinic / SCORE can be $0; private $200-500."),
    ("Indiana LLC", "Before taking money", f"={names['llc']}", False, "Inputs (roadmap §4)."),
    ("Domain + email for the link page", "Before beta", 20, True, "Estimate."),
]
r = 5
for i, (what, when, amt, is_input, src) in enumerate(cash, start=1):
    put(wsC, f"A{r}", i)
    put(wsC, f"B{r}", what).alignment = WRAP
    put(wsC, f"C{r}", when)
    if is_input:
        put(wsC, f"D{r}", amt, BLUE, USD, INPUT_FILL)
    else:
        put(wsC, f"D{r}", amt, BLACK, USD)
    put(wsC, f"E{r}", src).alignment = WRAP
    r += 1
c_last = r - 1
put(wsC, f"B{r}", "Subtotal", BOLD)
put(wsC, f"D{r}", f"=SUM(D5:D{c_last})", BLACK, USD, TOT_FILL)
c_sub = r
r += 1
put(wsC, f"B{r}", "Contingency, % of subtotal (re-orders, a dead camera, express shipping)")
put(wsC, f"C{r}", 0.10, BLUE, PCT, INPUT_FILL)
put(wsC, f"D{r}", f"=D{c_sub}*C{r}", BLACK, USD)
c_cont = r
r += 1
put(wsC, f"B{r}", "TOTAL PLANNED SPEND", BOLD)
put(wsC, f"D{r}", f"=D{c_sub}+D{c_cont}", BLACK, USD, TOT_FILL, bold=True)
c_tot = r
r += 1
put(wsC, f"B{r}", "Budget")
put(wsC, f"D{r}", f"={names['budget']}", BLACK, USD)
r += 1
put(wsC, f"B{r}", "LEFT OVER (negative = over budget)", BOLD)
put(wsC, f"D{r}", f"={names['budget']}-D{c_tot}", BLACK, USD, TOT_FILL, bold=True)
c_left = r
names["cash_total"] = ref(C, "D", c_tot)
names["cash_left"] = ref(C, "D", c_left)
r += 2
put(wsC, f"A{r}", "Not inside the $2,000: what the first 100 public units need", BOLD)
r += 1
put(wsC, f"B{r}", "100-unit batch (LCD version): variable cost × 100 + its one-time costs (samples, jig)")
put(wsC, f"D{r}", f"=('{U}'!E{varc_row}+'{V}'!D{vrows['d_lcd']})*'{U}'!E5+'{U}'!E{once_tot}", BLACK, USD0)
n_batch = r
r += 1
put(wsC, f"B{r}", "Selling legally: FCC SDoC, lawyer, LLC, trademark, insurance")
put(wsC, f"D{r}", f"={names['cert_total']}", BLACK, USD0)
n_cert = r
r += 1
put(wsC, f"B{r}", "Cash needed before the first public sale", BOLD)
put(wsC, f"D{r}", f"=D{n_batch}+D{n_cert}", BLACK, USD0, TOT_FILL, bold=True)
n_need = r
names["need100"] = ref(C, "D", r)
r += 1
put(wsC, f"B{r}", "Pre-orders at $225 (after card fees) that fund it")
put(wsC, f"D{r}", f"=ROUNDUP((D{n_need}-MAX(0,D{c_left}))/({names['hw_price']}*(1-{names['pay_pct']})-{names['pay_fixed']}),0)", BLACK, NUM)
names["preorders"] = ref(C, "D", r)
put(wsC, f"E{r}", "Uses any money left from the $2,000 first. Roadmap §6: take pre-orders before each step up.")

# ---------------------------------------------------------------- Checks
wsK = wb.create_sheet("Checks")
K = "Checks"
put(wsK, "A1", "Cross-checks against earlier files", TITLE)
header(wsK, 3, ["Check", "This model", "Earlier figure", "Difference", "Earlier source"], [60, 13, 13, 13, 70])
checks = [
    ("5 testers: cash for the batch minus the 2 owned shells' value", f"{names['batch_cash']['C']}-2*'{U}'!C9", 663.97, "AI_Calculator_Shopping_List.xlsx option C total ($663.97)"),
    ("LCD version cost per unit at 100 (before selling-legally)", names["V_lcd"]["E"], LCD100_ALIBABA, "Alibaba_Bulk_Sourcing.xlsx 'v15-LCD cost by batch' total @100"),
    ("Cost per unit at 10 (before selling-legally)", names["cpu"]["D"], 92, "Alibaba_Bulk_Sourcing.xlsx 'Cost by batch' (≈ $92)"),
    ("Cost per unit at 100 (before selling-legally)", names["cpu"]["E"], 46, "Alibaba_Bulk_Sourcing.xlsx (≈ $46)"),
    ("Cost per unit at 1,000 (before selling-legally)", names["cpu"]["F"], 34, "Alibaba_Bulk_Sourcing.xlsx (≈ $34)"),
    ("Opus 5.5 cost per solve, no caching (roadmap §3.3)", f"(({names['in_tokens']})*4+2100*20)/1000000", 0.052, "LAUNCH_ROADMAP.md §3.3 (~$0.052)"),
    ("Certification etc. one-time", names["cert_total"], 3500, "LAUNCH_ROADMAP.md §3 ('about $3.5k')"),
]
for i, (lab, f_, earlier, src) in enumerate(checks, start=4):
    put(wsK, f"A{i}", lab)
    put(wsK, f"B{i}", f"={f_}", BLACK, USD if earlier > 1 else '"$"0.0000')
    put(wsK, f"C{i}", earlier, BLUE, USD if earlier > 1 else '"$"0.0000', INPUT_FILL)
    put(wsK, f"D{i}", f"=B{i}-C{i}", BLACK, USD if earlier > 1 else '"$"0.0000')
    put(wsK, f"E{i}", src)
put(wsK, "A12", "Differences come from the tester-only rows, rounding in the earlier sheet, and the cache saving on the system prompt.", NOTE)

# ---------------------------------------------------------------- Summary (first)
wsS = wb.create_sheet("Summary", 0)
put(wsS, "A1", "AI Calculator cost model: headline numbers", TITLE)
put(wsS, "A2", "All values are formulas pointing at the other sheets. Edit blue cells on 'Inputs', 'Unit cost' and 'Cash budget'.", NOTE)
header(wsS, 4, ["", "5 testers", "10 units", "100 units", "1,000 units"], [62, 14, 14, 14, 14])
summ = [
    ("Cost per unit (parts, freight, scrap, batch one-time)", lambda uc, mc: f"={names['cpu'][uc]}", USD),
    ("Cost per unit incl. FCC / lawyer / LLC / trademark / insurance", lambda uc, mc: f"={names['cpu_all'][uc]}", USD),
    ("Hardware profit per $225 unit", lambda uc, mc: f"='{M}'!{mc}13", USD),
    ("Hardware margin %", lambda uc, mc: f"='{M}'!{mc}14", PCT),
    ("Subscription profit per paying user per month", lambda uc, mc: f"='{M}'!{mc}20", USD),
    ("Subscription margin %", lambda uc, mc: f"='{M}'!{mc}21", PCT),
    ("Total profit per calculator (hardware + 12 months' subscription)", lambda uc, mc: f"='{M}'!{mc}27", USD),
    ("Units to cover this batch's one-time costs", lambda uc, mc: f"='{M}'!{mc}32", NUM),
]
for i, (lab, fn, fmt) in enumerate(summ, start=5):
    put(wsS, f"A{i}", lab)
    for uc, mc, col in zip(VOLS, MC, ["B", "C", "D", "E"]):
        put(wsS, f"{col}{i}", fn(uc, mc), BLACK, fmt)
r = 5 + len(summ) + 1
put(wsS, f"A{r}", "Versions (10 / 100 / 1,000 only; the 5 testers are e-paper)", BOLD)
for lab, key in [("LCD version: cost per unit incl. selling-legally", "V_lcd_all"),
                 ("LCD version: hardware profit per $225 unit", "V_lcd_profit"),
                 ("E-paper + 1,200-1,500 mAh: cost per unit incl. selling-legally", "V_ep_big_all"),
                 ("LCD + 1,200-1,500 mAh: cost per unit incl. selling-legally", "V_lcd_big_all"),
                 ("LCD + 1,200-1,500 mAh: hardware profit per $225 unit", "V_lcd_big_profit")]:
    r += 1
    put(wsS, f"A{r}", lab)
    put(wsS, f"B{r}", "n/a", NOTE)
    for uc, col in zip(["D", "E", "F"], ["C", "D", "E"]):
        put(wsS, f"{col}{r}", f"={names[key][uc]}", BLACK, USD)
r += 2
put(wsS, f"A{r}", "Claude API", BOLD)
singles = [
    ("Model", f"={names['model']}", None),
    ("Cost per solve", f"={names['api_solve']}", '"$"0.0000'),
    ("API cost per subscriber per month (typical use)", f"={names['api_month']}", USD),
    ("API cost per month for a user at the 500 cap", f"={names['api_cap']}", USD),
    ("Solves per day that $15 pays for (break-even)", f"={names['api_be_day']}", NUM2),
    ("Suggested fair-use cap (solves/month) for break-even", f"={names['api_be_cap']}", NUM),
    ("Same API cost per month on Sonnet 5.5 (typical)", f"='{A}'!D{names['scen_first']+1}", USD),
    ("Same on Sonnet 5.5 for a user at the 500 cap", f"='{A}'!F{names['scen_first']+1}", USD),
]
for lab, f_, fmt in singles:
    r += 1
    put(wsS, f"A{r}", lab)
    put(wsS, f"B{r}", f_, BLACK, fmt)
r += 2
put(wsS, f"A{r}", "The $2,000", BOLD)
for lab, f_, fmt in [("Planned spend (5 testers → v15 order + panels → v15-LCD pilot)", f"={names['cash_total']}", USD),
                     ("Left over", f"={names['cash_left']}", USD),
                     ("Cash needed before the first 100 public units", f"={names['need100']}", USD0),
                     ("Pre-orders at $225 that would fund it", f"={names['preorders']}", NUM)]:
    r += 1
    put(wsS, f"A{r}", lab)
    put(wsS, f"B{r}", f_, BLACK, fmt)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = True
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)

wb.save(OUT)
print("saved", OUT)
